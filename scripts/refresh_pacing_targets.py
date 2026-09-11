"""Regenerate config/dtc_pacing_targets.json from an .xlsx export of the ecomm
team's pacing sheet ("2026 Daily Pacing & Performance D2C | Biom").

    uv run python scripts/refresh_pacing_targets.py <export.xlsx> [--out PATH]
        [--file-id ID] [--url URL] [--check]

How to get the export: File > Download > Microsoft Excel (.xlsx) in Google
Sheets, or the Drive connector's xlsx export in a Claude Code session. The
sheet is updated monthly (a new "<Month> <Year>" tab appears at the start of
each month), so run this once a month, then commit the JSON.

What is read, per monthly tab (`^<Month> <YYYY>$`; tabs like "April 2026 V2"
and the source tabs are skipped and listed):

  * the header row (the row whose first cell is "Week" and third is "Date");
  * one row per calendar day, from the row after the header while column C
    holds a date;
  * the summary block below the days: the "Total" row's Forecast / Forecasted
    Spend / Forecasted NCs cells become `totals`.

Columns are matched on whitespace-normalised header text via
`bpd_mcp.pacing_targets.FIELD_MAP`. A column the tab does not carry is simply
absent for that month (the tool treats a missing series as "no target"); a
column the sheet renames therefore shows up as a missing series, never as the
wrong series. Non-numeric cells ("-", "#REF!", blanks) become null.

`--check` re-reads the export and exits 1 if the JSON on disk would change —
for a Routine that wants to know when the sheet has a new tab.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from bpd_mcp.pacing_targets import (
    DEFAULT_TARGETS_PATH,
    FIELD_MAP,
    PACED_FIELDS,
    REFRESH_COMMAND,
    SCHEMA_VERSION,
    coerce_number,
    parse_targets,
)

SHEET_TITLE = "2026 Daily Pacing & Performance D2C | Biom"
DEFAULT_FILE_ID = "18YnVLGUaVP5gNVJgVPWxuiC9OPYJbT6S6_A8Os1SVM4"

_MONTHS = (
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)
_TAB_RE = re.compile(r"^(" + "|".join(_MONTHS) + r") (\d{4})$")


def _norm(v: Any) -> str | None:
    if v is None:
        return None
    s = " ".join(str(v).split())
    return s or None


def _as_date(v: Any) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _find_header_row(ws: Any) -> int | None:
    for r in range(1, min(ws.max_row, 20) + 1):
        if _norm(ws.cell(r, 1).value) == "Week" and _norm(ws.cell(r, 3).value) == "Date":
            return r
    return None


def read_month_tab(ws: Any) -> dict[str, Any] | None:
    """One tab -> {"tab", "days", "totals"} or None when the tab has no pacing grid."""
    header_row = _find_header_row(ws)
    if header_row is None:
        return None
    header: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = _norm(ws.cell(header_row, c).value)
        if h and h not in header:
            header[h] = c
    col_for = {
        fld: header[sheet_col] for fld, sheet_col in FIELD_MAP.items() if sheet_col in header
    }

    days: dict[str, dict[str, float | None]] = {}
    r = header_row + 1
    while r <= ws.max_row:
        d = _as_date(ws.cell(r, 3).value)
        if d is None:
            break
        days[d.isoformat()] = {
            fld: coerce_number(ws.cell(r, c).value) for fld, c in col_for.items()
        }
        r += 1
    if not days:
        return None

    totals: dict[str, float | None] = dict.fromkeys(PACED_FIELDS)
    for rr in range(r, min(r + 15, ws.max_row) + 1):
        if _norm(ws.cell(rr, 3).value) == "Total":
            for fld in PACED_FIELDS:
                c = col_for.get(fld)
                if c is not None:
                    totals[fld] = coerce_number(ws.cell(rr, c).value)
            break
    return {
        "tab": ws.title,
        "days": days,
        "totals": totals,
        "missing_fields": sorted(set(FIELD_MAP) - set(col_for)),
    }


def build_payload(xlsx: Path, *, file_id: str, url: str | None) -> tuple[dict[str, Any], list[str]]:
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # pivot-cache relationship noise from the export
        wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=False)
    months: dict[str, Any] = {}
    skipped: list[str] = []
    notes: list[str] = []
    for name in wb.sheetnames:
        m = _TAB_RE.match(name.strip())
        if not m:
            skipped.append(name)
            continue
        ym = f"{m[2]}-{_MONTHS.index(m[1]) + 1:02d}"
        tab = read_month_tab(wb[name])
        if tab is None:
            skipped.append(name)
            continue
        if tab["missing_fields"]:
            notes.append(f"{name}: no column for {tab['missing_fields']}")
        months[ym] = {"tab": tab["tab"], "totals": tab["totals"], "days": tab["days"]}
    payload = {
        "schema_version": SCHEMA_VERSION,
        "source": {
            "title": SHEET_TITLE,
            "file_id": file_id,
            "url": url or f"https://docs.google.com/spreadsheets/d/{file_id}/edit",
            "export_file": xlsx.name,
            "export_modified_at": dt.datetime.fromtimestamp(
                xlsx.stat().st_mtime, tz=dt.UTC
            ).isoformat(timespec="seconds"),
            "refreshed_at": dt.datetime.now(tz=dt.UTC).isoformat(timespec="seconds"),
            "refresh_command": REFRESH_COMMAND,
            "field_map": FIELD_MAP,
            "skipped_tabs": skipped,
            "notes": notes,
        },
        "months": dict(sorted(months.items())),
    }
    return payload, notes


def _stable(payload: dict[str, Any]) -> dict[str, Any]:
    """The payload minus the timestamps, for --check comparisons."""
    p = json.loads(json.dumps(payload))
    p["source"].pop("refreshed_at", None)
    p["source"].pop("export_modified_at", None)
    p["source"].pop("export_file", None)
    return p


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("xlsx", type=Path, help=".xlsx export of the pacing sheet")
    ap.add_argument("--out", type=Path, default=DEFAULT_TARGETS_PATH)
    ap.add_argument("--file-id", default=DEFAULT_FILE_ID)
    ap.add_argument("--url", default=None)
    ap.add_argument("--check", action="store_true", help="exit 1 if the JSON on disk would change")
    args = ap.parse_args(argv)

    payload, notes = build_payload(args.xlsx, file_id=args.file_id, url=args.url)
    parsed = parse_targets(payload)  # validates the shape before anything is written
    for ym in parsed.available_months:
        mt = parsed.months[ym]
        stated = mt.totals.get("forecast_demand")
        summed = mt.series_sum("forecast_demand") or 0.0
        flag = ""
        if stated is not None and summed and abs(stated - summed) > 0.005 * max(abs(stated), 1):
            flag = f"  ** Total row {stated:,.2f} != day sum {summed:,.2f}"
        print(f"{ym}  {mt.tab:16s} days={len(mt.days):2d}  forecast_demand={summed:>12,.2f}{flag}")
    for n in notes:
        print("note:", n)
    if payload["source"]["skipped_tabs"]:
        print("skipped tabs:", ", ".join(payload["source"]["skipped_tabs"]))

    if args.check:
        if not args.out.exists():
            print(f"--check: {args.out} does not exist")
            return 1
        current = json.loads(args.out.read_text())
        if _stable(current) != _stable(payload):
            print(f"--check: {args.out} is out of date with {args.xlsx}")
            return 1
        print(f"--check: {args.out} is current")
        return 0

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, indent=1, sort_keys=False) + "\n")
    print(f"wrote {args.out} ({len(parsed.months)} month(s))")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
